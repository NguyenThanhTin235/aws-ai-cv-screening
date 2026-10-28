import os
import re
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt, Inches, RGBColor

CVS_DIR = os.path.join(os.path.dirname(__file__), '..', 'samples', 'cvs')

def md_to_pdf(md_path, pdf_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = SimpleDocTemplate(pdf_path, pagesize=letter,
                            rightMargin=40, leftMargin=40,
                            topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155')
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        leftIndent=15,
        textColor=colors.HexColor('#334155')
    )

    story = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 4))
            continue
        if stripped.startswith('# '):
            text = stripped[2:].replace('&', '&amp;')
            story.append(Paragraph(f"<b>{text}</b>", title_style))
        elif stripped.startswith('## '):
            text = stripped[3:].replace('&', '&amp;')
            story.append(Paragraph(f"<b>{text}</b>", h2_style))
        elif stripped.startswith('### '):
            text = stripped[4:].replace('&', '&amp;')
            story.append(Paragraph(f"<b>{text}</b>", body_style))
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:].replace('&', '&amp;')
            text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
            story.append(Paragraph(f"• {text}", bullet_style))
        elif stripped.startswith('---'):
            story.append(Spacer(1, 6))
        else:
            text = stripped.replace('&', '&amp;')
            text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
            story.append(Paragraph(text, body_style))

    doc.build(story)
    print(f"Generated PDF: {pdf_path}")

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith('# '):
            p = doc.add_heading(level=1)
            run = p.add_run(stripped[2:])
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        elif stripped.startswith('## '):
            p = doc.add_heading(level=2)
            run = p.add_run(stripped[3:])
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.size = Pt(10)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        elif stripped.startswith('---'):
            continue
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

def main():
    cv_files = [f for f in os.listdir(CVS_DIR) if f.endswith('.md')]
    for cv in cv_files:
        base_name = os.path.splitext(cv)[0]
        md_path = os.path.join(CVS_DIR, cv)
        
        # Tạo PDF cho tất cả CV
        pdf_path = os.path.join(CVS_DIR, f"{base_name}.pdf")
        md_to_pdf(md_path, pdf_path)
        
        # Tạo DOCX cho các CV test Word (CV 02, 04, 05)
        if any(k in base_name for k in ['backend', 'sales', 'pii']):
            docx_path = os.path.join(CVS_DIR, f"{base_name}.docx")
            md_to_docx(md_path, docx_path)

if __name__ == '__main__':
    main()
